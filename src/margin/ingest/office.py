"""Word and PowerPoint files.

Word: headings come from paragraph styles ("Heading 1", "Title"), which is far
more reliable than guessing from text. Paragraphs and tables are read in
document order by walking the body XML, because python-docx lists them
separately. Word files have no fixed pages, so blocks carry ``page=None``.

PowerPoint: each slide is a page, its title a heading, and speaker notes are
kept — lecturers often put the explanation there rather than on the slide.
"""

from __future__ import annotations

from pathlib import Path

from margin.ingest.types import Block, Document

TITLE_STYLES = {"title": 1}


def _heading_level(style_name: str) -> int | None:
    name = style_name.lower().strip()
    if name in TITLE_STYLES:
        return TITLE_STYLES[name]
    if name.startswith("heading "):
        suffix = name.removeprefix("heading ").strip()
        return int(suffix) if suffix.isdigit() else None
    return None


def _table_text(rows: list[list[str]]) -> str:
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows if any(c.strip() for c in row))


def read_docx(path: Path, doc_id: str) -> Document:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    blocks: list[Block] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            level = _heading_level(para.style.name if para.style is not None else "")
            blocks.append(Block(text, None, "heading", level) if level else Block(text, None))
        elif tag == "tbl":
            table = Table(child, document)
            text = _table_text([[cell.text for cell in row.cells] for row in table.rows])
            if text:
                blocks.append(Block(text, None, "table"))
    title = (document.core_properties.title or "").strip() or path.stem
    return Document(doc_id, str(path), "docx", title, 0, tuple(blocks))


def read_pptx(path: Path, doc_id: str) -> Document:
    from pptx import Presentation

    prs = Presentation(str(path))
    blocks: list[Block] = []
    for number, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        # python-pptx builds a new proxy object on every access, so compare ids, not identity
        title_id = title_shape.shape_id if title_shape is not None else None
        blocks.append(Block(title or f"Slide {number}", number, "heading", 1))
        for shape in slide.shapes:
            if shape.shape_id == title_id:
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                blocks.append(Block(shape.text_frame.text.strip(), number))
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                if text := _table_text(rows):
                    blocks.append(Block(text, number, "table"))
        if slide.has_notes_slide and (notes := slide.notes_slide.notes_text_frame.text.strip()):
            blocks.append(Block(notes, number, "notes"))
    title = (prs.core_properties.title or "").strip() or path.stem
    return Document(doc_id, str(path), "pptx", title, len(prs.slides), tuple(blocks))
