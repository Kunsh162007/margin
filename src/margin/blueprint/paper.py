"""Read a past exam paper into numbered questions with marks and a year.

Papers are read line by line from the raw page text. The book reader merges
lines into paragraphs, which is right for prose and wrong here: a question
number is only recognisable at the start of a line.

Recognised: ``Q3.``, ``Question 3``, ``3.``, ``3)``, sub-parts ``(a)`` or
``(ii)`` on their own line or straight after the number, and marks written as
``[5 marks]``, ``(5 marks)``, ``[5]`` or ``(5)`` at the end of a question. A
question stem is carried into each of its parts so that "(b) Find the speed"
is searched together with the situation it refers to.
"""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

MIN_QUESTION_CHARS = 12
MIN_TEXT_PER_PAGE = 40
_QUESTION_HEAD = re.compile(
    r"^(?:Q(?:uestion)?\s*\.?\s*(\d{1,2})\s*[.):]?|(\d{1,2})\s*[.):])\s*(?:\(\s*([a-h]|[ivx]{1,4})\s*\))?\s*",
    re.IGNORECASE,
)
_SUBPART_HEAD = re.compile(r"^\(\s*([a-h]|[ivx]{1,4})\s*\)\s*", re.IGNORECASE)
_MARKS = re.compile(r"[\[(]\s*(\d{1,3})\s*(?:marks?|m)?\s*[\])]\s*$", re.IGNORECASE)
_YEAR = re.compile(r"\b(19[89]\d|20[0-4]\d)\b")
_NOISE = re.compile(r"^(?:page\s+\d+(?:\s+of\s+\d+)?|turn over|p\.?\s*t\.?\s*o\.?|end of (?:paper|examination|exam))$", re.IGNORECASE)


@dataclass(frozen=True)
class PaperQuestion:
    number: str  # "3", "3(b)"
    text: str
    marks: int | None
    year: int | None
    source: str = ""

    @property
    def key(self) -> str:
        return f"{self.source}:{self.number}" if self.source else self.number


def detect_year(text: str) -> int | None:
    """The year printed in the paper's header, if any."""
    years = _YEAR.findall(text[:800])
    return int(Counter(years).most_common(1)[0][0]) if years else None


def _split_marks(body: str) -> tuple[str, int | None]:
    m = _MARKS.search(body)
    return (body[: m.start()].rstrip(), int(m.group(1))) if m else (body, None)


def parse_paper(text: str, source: str = "") -> list[PaperQuestion]:
    year = detect_year(text)
    questions: list[PaperQuestion] = []
    number: int | None = None
    part: str | None = None
    stem = ""
    buffer: list[str] = []

    def flush() -> None:
        if number is None or not buffer:
            return
        body, marks = _split_marks(" ".join(buffer).strip())
        if len(body) < MIN_QUESTION_CHARS:
            return
        label = f"{number}({part})" if part else str(number)
        questions.append(PaperQuestion(label, f"{stem} {body}".strip() if part else body, marks, year, source))

    for raw in text.splitlines():
        line = raw.strip()
        if not line or _NOISE.match(line):
            continue
        head = _QUESTION_HEAD.match(line)
        if head:
            candidate = int(head.group(1) or head.group(2))
            if number is None or candidate in (number, number + 1):
                flush()
                if candidate != number:
                    stem = ""
                number, part, buffer = candidate, (head.group(3) or "").lower() or None, [line[head.end() :]]
                continue
        sub = _SUBPART_HEAD.match(line)
        if sub and number is not None:
            if part is None:
                stem = _split_marks(" ".join(buffer).strip())[0]
            else:
                flush()
            part, buffer = sub.group(1).lower(), [line[sub.end() :]]
            continue
        if number is not None:
            buffer.append(line)
    flush()
    return questions


def paper_text(path: Path) -> str:
    """Raw text of a paper file, keeping line breaks. Scanned pages and photos go through OCR."""
    from margin.ingest.detect import sniff

    kind = sniff(path)
    if kind == "pdf":
        return _pdf_text(path)
    if kind == "docx":
        import docx

        return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    if kind in ("png", "jpeg", "webp"):
        from margin.ingest.images import OcrEngine, load_image

        return OcrEngine().read(load_image(path)).text
    if kind == "text":
        return path.read_text(encoding="utf-8", errors="replace")
    from margin.ingest.pipeline import read_document

    return "\n".join(b.text for b in read_document(path).blocks)


def _pdf_text(path: Path) -> str:
    import pypdfium2 as pdfium

    from margin.ingest.images import OcrEngine

    pdf = pdfium.PdfDocument(str(path))
    ocr: OcrEngine | None = None
    pages: list[str] = []
    try:
        for index in range(len(pdf)):
            page = pdf[index]
            textpage = page.get_textpage()
            text = textpage.get_text_bounded()
            textpage.close()
            if len(text.strip()) < MIN_TEXT_PER_PAGE:
                ocr = ocr or OcrEngine()
                text = ocr.read(page.render(scale=2.0).to_pil()).text
            page.close()
            pages.append(text)
    finally:
        pdf.close()
    return "\n".join(pages)
