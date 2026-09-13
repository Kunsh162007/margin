"""Run the whole study flow through the real interface, with the real model and no network.

    uv run python -m evals.ui_flow_eval

Phase 7's criterion is "the whole flow runs from the UI". This drives the Textual
app headlessly, pressing the same buttons a student would, with the default model
(Gemma 4 E4B) and the real search models, while every connection that leaves the
machine is refused. The library is a fresh copy of the University Physics eval
workspace, so this run never changes what the other evals read.

Steps, in a student's order: add a Word file, read a past paper, write notes for
one section with a diagram, export HTML and an Anki deck, write questions, mark
the model answer to one of them, and ask the agent for a comparison table.
Recorded per step: whether it did what it should, and seconds. Refused network
attempts are counted, so a library that catches the error and carries on shows up.
"""

from __future__ import annotations

import asyncio
import json
import shutil
import sys
import time
from datetime import datetime
from pathlib import Path

import docx
from textual.widgets import Checkbox, DataTable, Input, OptionList, ProgressBar, RichLog, Static, TabbedContent, TextArea

from evals.netguard import blocked_network
from margin import config
from margin.export.assets import ensure_mermaid, mermaid_script
from margin.retrieve.embed import Embedder
from margin.runtime.session import ModelSession, server_config
from margin.store.db import Workspace
from margin.study import Study
from margin.tui.app import MarginApp

SID = "6.2"
PAPER_SECTIONS = ("5.3", "6.2", "10.6")
GOLD = Path(__file__).parent / "datasets" / "retrieval_gold.jsonl"
OUT_DIR = Path(__file__).parent / "results" / "ui"
ASK = f"Make a table comparing static and kinetic friction, using section {SID}."


def _paper(path: Path) -> int:
    gold = [json.loads(line) for line in GOLD.read_text(encoding="utf-8").splitlines() if line.strip()]
    picked = [g["question"] for sid in PAPER_SECTIONS for g in [g for g in gold if g["gold_section"] == sid][:2]]
    lines = ["PHYSICS EXAMINATION 2025", ""] + [f"Q{i + 1}. {q} [{3 + i % 3} marks]" for i, q in enumerate(picked)]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return len(picked)


def _word_file(path: Path) -> Path:
    document = docx.Document()
    document.add_heading("Friction revision sheet", level=1)
    document.add_paragraph("Static friction holds a resting object in place up to a maximum value. Once the object slides, "
                           "kinetic friction acts instead, and it is usually smaller than the maximum static friction.")
    document.save(str(path))
    return path


def _saved_visuals(workspace: Path, sid: str) -> list[tuple[str, str]]:
    with Workspace.open(workspace) as ws:
        found = ws.find_section(sid)
        saved = ws.load_notes(found[0], sid) if found else None
    return [(v["kind"], v["format"]) for v in json.loads(saved["visuals_json"])] if saved else []


async def _open(app: MarginApp, pilot, tab: str) -> None:
    """Click the tab in the tab bar, as a student would, and confirm it opened."""
    target = next(t for t in app.screen.query("ContentTab") if t.id == f"--content-tab-{tab}")
    await pilot.click(target)
    await pilot.pause()
    if app.query_one("#tabs", TabbedContent).active != tab:
        raise RuntimeError(f"the {tab} tab did not open")


async def _flow(app: MarginApp, study: Study, work_dir: Path, record) -> None:
    shots = work_dir / "screenshots"  # one SVG per step, for the README and the demo
    async with app.run_test(size=(140, 42)) as pilot:

        async def press(selector: str) -> bool:
            target = app.query_one(selector)
            if target.region.area == 0:  # not on screen: a click would land on whatever is at the origin
                return False
            clicked = await pilot.click(target)
            await app.workers.wait_for_complete()
            await pilot.pause()
            return clicked

        async def step(name: str, action, check) -> None:
            while len(app.screen_stack) > 1:  # close anything a stray click opened, e.g. the command palette
                app.pop_screen()
            await pilot.pause()
            started = time.perf_counter()
            clicked = await action()
            outcome = check() if clicked else (False, "")
            ok, detail = outcome if isinstance(outcome, tuple) else (bool(outcome), "")
            status = str(app.query_one("#status", Static).render()) if clicked else "click missed: the control was not on screen"
            record(name, bool(clicked) and ok, time.perf_counter() - started, status, detail)
            shots.mkdir(exist_ok=True)
            app.save_screenshot(filename=f"{len(list(shots.glob('*.svg'))) + 1:02d}-{name.replace(' ', '-')}.svg", path=str(shots))

        documents = app.query_one("#documents", DataTable)
        before = documents.row_count
        word = _word_file(work_dir / "friction-revision.docx")
        app.query_one("#add-paths", Input).value = str(word)
        await step("add a Word file", lambda: press("#add"), lambda: documents.row_count == before + 1)

        await _open(app, pilot, "blueprint")
        paper = work_dir / "physics-2025.txt"
        asked = _paper(paper)
        app.query_one("#papers", Input).value = str(paper)
        await step("read a past paper", lambda: press("#read-papers"),
                   lambda: study.blueprint is not None and study.blueprint.questions == asked and app.query_one("#chapters", DataTable).row_count > 0)

        await _open(app, pilot, "notes")
        app.query_one("#notes-scope", Input).value = SID
        app.query_one("#visuals", Checkbox).value = True
        bar = app.query_one("#notes-progress", ProgressBar)
        def notes_written():
            visuals = _saved_visuals(study.workspace, SID)
            done = bar.total is not None and bar.total >= 1 and bar.progress == bar.total
            return done and bool(visuals), f"visuals saved: {visuals}"

        await step("write notes with a diagram", lambda: press("#write-notes"), notes_written)
        html_path = study.export_dir / f"margin-{SID.replace('.', '-')}.html"

        def html_exported():
            if not html_path.exists():
                return False, "no file written"
            page = html_path.read_text(encoding="utf-8")
            visuals = _saved_visuals(study.workspace, SID)
            diagrams = sum(1 for _, fmt in visuals if fmt == "mermaid")
            tables = sum(1 for kind, fmt in visuals if fmt != "mermaid" and kind == "table")
            ok = (page.count('class="mermaid"') >= diagrams and (diagrams == 0 or "mermaid.initialize" in page)
                  and page.count("<table>") >= tables)
            return ok, f"{html_path.stat().st_size:,} bytes; diagrams {page.count('class=\"mermaid\"')} (saved {diagrams}); tables {page.count('<table>')} (saved {tables})"

        await step("export HTML", lambda: press("#export-html"), html_exported)
        await step("export an Anki deck", lambda: press("#export-anki"), lambda: (study.export_dir / f"margin-{SID.replace('.', '-')}.apkg").exists())

        await _open(app, pilot, "practice")
        app.query_one("#practice-scope", Input).value = SID
        app.query_one("#count", Input).value = "3"
        app.query_one("#marks", Input).value = "2"
        await step("write questions", lambda: press("#generate"), lambda: app.query_one("#questions", OptionList).option_count >= 1)
        if study.questions:
            first = study.questions[0]
            app.query_one("#answer", TextArea).load_text(first.answer)
            await step("mark the model answer", lambda: press("#mark"),
                       lambda: f"{first.marks}/{first.marks} marks" in str(app.query_one("#feedback", Static).render()))

        await _open(app, pilot, "ask")
        box = app.query_one("#ask-input", Input)
        box.value = ASK
        box.focus()

        async def submit() -> bool:
            await pilot.press("enter")
            await app.workers.wait_for_complete()
            await pilot.pause()
            return True

        log = app.query_one("#ask-log", RichLog)

        def table_made():
            transcript = "\n".join(line.text.rstrip() for line in log.lines)
            return "make_table done" in transcript, transcript[-600:]

        await step("ask for a comparison table", submit, table_made)


def main() -> int:
    paths = config.paths().ensure()
    source = paths.workspaces_dir / "eval-uphys1.db"
    if not source.exists():
        print(f"no workspace at {source}; run evals.retrieval_eval first")
        return 1
    ensure_mermaid()  # fetched before the network is switched off, as margin setup does
    work_dir = paths.home / "ui-flow"
    shutil.rmtree(work_dir, ignore_errors=True)
    work_dir.mkdir(parents=True)
    workspace = work_dir / "library.db"
    shutil.copy2(source, workspace)
    with Workspace.open(workspace) as ws:
        ws.clear_notes([SID])

    spec, backend, cfg = server_config()
    session = ModelSession(spec, backend, cfg, paths.logs_dir / f"ui-flow-eval-{spec.id}.log")
    study = Study(workspace, work_dir / "exports", session.client, Embedder(), mermaid_script)
    app = MarginApp(study, model_label=f"{spec.id} on {backend.upper()}")
    steps: list[dict[str, object]] = []

    def record(name: str, ok: bool, seconds: float, status: str, detail: str = "") -> None:
        steps.append({"step": name, "ok": ok, "seconds": round(seconds, 1), "status": status, "detail": detail})
        print(f"{'ok  ' if ok else 'FAIL'} {name}: {seconds:.1f} s — {status}" + (f"\n     {detail}" if detail else ""), flush=True)

    started = time.perf_counter()
    attempts: list[str] = []
    try:
        with blocked_network(attempts):
            asyncio.run(_flow(app, study, work_dir, record))
    finally:
        session.close()

    summary = {
        "model": spec.id, "backend": backend, "steps": steps, "all_ok": bool(steps) and all(s["ok"] for s in steps),
        "steps_ok": sum(1 for s in steps if s["ok"]), "steps_run": len(steps), "total_seconds": round(time.perf_counter() - started, 1),
        "network_attempts_refused": len(attempts), "refused_targets": sorted(set(attempts))[:10],
        "questions_accepted": len(study.questions), "run_at": datetime.now().strftime("%Y%m%d_%H%M%S"),
    }
    print(json.dumps(summary, indent=2))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / "uphys1.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    return 0 if summary["all_ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
